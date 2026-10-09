import streamlit as st
import google.generativeai as genai
import pypdf
import docx
import os

# Page Config
st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="centered")
st.title("📄 AI Resume ATS Checker")
st.write("Apna Resume upload karo aur Job Description paste karo.")

# API Key - Streamlit Secrets se lega
try:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
    genai.configure(api_key=GOOGLE_API_KEY)
except:
    st.error("API Key nahi mili! Streamlit > Advanced Settings > Secrets me GOOGLE_API_KEY add karo.")
    st.stop()

# Function to read PDF/DOCX
def extract_text(uploaded_file):
    if uploaded_file.type == "application/pdf":
        reader = pypdf.PdfReader(uploaded_file)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
        return text
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(uploaded_file)
        return "\n".join([p.text for p in doc.paragraphs])
    else:
        return uploaded_file.read().decode("utf-8")

# Inputs
job_desc = st.text_area("Paste Job Description here:", height=200)
resume_file = st.file_uploader("Upload Resume (PDF / DOCX)", type=["pdf", "docx"])

if st.button("Check ATS Score"):
    if not job_desc or not resume_file:
        st.warning("Pehle Job Description aur Resume dono do.")
    else:
        with st.spinner("AI Check kar raha hai..."):
            resume_text = extract_text(resume_file)
            
            model = genai.GenerativeModel("gemini-1.5-flash")
            
            prompt = f"""
            You are an expert ATS (Applicant Tracking System).
            Compare this Resume with the Job Description.
            
            Job Description:
            {job_desc}
            
            Resume:
            {resume_text}
            
            Give output in this format:
            1. ATS Match Score in % (e.g. 75%)
            2. Missing Keywords
            3. Strengths
            4. Improvements / Suggestions to improve resume
            5. Final Verdict
            
            Keep language simple and in Roman Urdu / English mix.
            """
            
            response = model.generate_content(prompt)
            st.success("Result Ready!")
            st.markdown(response.text)import streamlit as st
import google.generativeai as genai
import pypdf
import docx
import os

# Page Config
st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="centered")
st.title("📄 AI Resume ATS Checker")
st.write("Apna Resume upload karo aur Job Description paste karo.")

# API Key - Streamlit Secrets se lega
try:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
    genai.configure(api_key=GOOGLE_API_KEY)
except:
    st.error("API Key nahi mili! Streamlit > Advanced Settings > Secrets me GOOGLE_API_KEY add karo.")
    st.stop()

# Function to read PDF/DOCX
def extract_text(uploaded_file):
    if uploaded_file.type == "application/pdf":
        reader = pypdf.PdfReader(uploaded_file)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
        return text
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(uploaded_file)
        return "\n".join([p.text for p in doc.paragraphs])
    else:
        return uploaded_file.read().decode("utf-8")

# Inputs
job_desc = st.text_area("Paste Job Description here:", height=200)
resume_file = st.file_uploader("Upload Resume (PDF / DOCX)", type=["pdf", "docx"])

if st.button("Check ATS Score"):
    if not job_desc or not resume_file:
        st.warning("Pehle Job Description aur Resume dono do.")
    else:
        with st.spinner("AI Check kar raha hai..."):
            resume_text = extract_text(resume_file)
            
            model = genai.GenerativeModel("gemini-1.5-flash")
            
            prompt = f"""
            You are an expert ATS (Applicant Tracking System).
            Compare this Resume with the Job Description.
            
            Job Description:
            {job_desc}
            
            Resume:
            {resume_text}
            
            Give output in this format:
            1. ATS Match Score in % (e.g. 75%)
            2. Missing Keywords
            3. Strengths
            4. Improvements / Suggestions to improve resume
            5. Final Verdict
            
            Keep language simple and in Roman Urdu / English mix.
            """
            
            response = model.generate_content(prompt)
            st.success("Result Ready!")
            st.markdown(response.text)import streamlit as st
import google.generativeai as genai
import pypdf
import docx
import os

# Page Config
st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="centered")
st.title("📄 AI Resume ATS Checker")
st.write("Apna Resume upload karo aur Job Description paste karo.")

# API Key - Streamlit Secrets se lega
try:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
    genai.configure(api_key=GOOGLE_API_KEY)
except:
    st.error("API Key nahi mili! Streamlit > Advanced Settings > Secrets me GOOGLE_API_KEY add karo.")
    st.stop()

# Function to read PDF/DOCX
def extract_text(uploaded_file):
    if uploaded_file.type == "application/pdf":
        reader = pypdf.PdfReader(uploaded_file)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
        return text
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(uploaded_file)
        return "\n".join([p.text for p in doc.paragraphs])
    else:
        return uploaded_file.read().decode("utf-8")

# Inputs
job_desc = st.text_area("Paste Job Description here:", height=200)
resume_file = st.file_uploader("Upload Resume (PDF / DOCX)", type=["pdf", "docx"])

if st.button("Check ATS Score"):
    if not job_desc or not resume_file:
        st.warning("Pehle Job Description aur Resume dono do.")
    else:
        with st.spinner("AI Check kar raha hai..."):
            resume_text = extract_text(resume_file)
            
            model = genai.GenerativeModel("gemini-1.5-flash")
            
            prompt = f"""
            You are an expert ATS (Applicant Tracking System).
            Compare this Resume with the Job Description.
            
            Job Description:
            {job_desc}
            
            Resume:
            {resume_text}
            
            Give output in this format:
            1. ATS Match Score in % (e.g. 75%)
            2. Missing Keywords
            3. Strengths
            4. Improvements / Suggestions to improve resume
            5. Final Verdict
            
            Keep language simple and in Roman Urdu / English mix.
            """
            
            response = model.generate_content(prompt)
            st.success("Result Ready!")
            st.markdown(response.text)import streamlit as st
import google.generativeai as genai
import pypdf
import docx
import os

# Page Config
st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="centered")
st.title("📄 AI Resume ATS Checker")
st.write("Apna Resume upload karo aur Job Description paste karo.")

# API Key - Streamlit Secrets se lega
try:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
    genai.configure(api_key=GOOGLE_API_KEY)
except:
    st.error("API Key nahi mili! Streamlit > Advanced Settings > Secrets me GOOGLE_API_KEY add karo.")
    st.stop()

# Function to read PDF/DOCX
def extract_text(uploaded_file):
    if uploaded_file.type == "application/pdf":
        reader = pypdf.PdfReader(uploaded_file)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
        return text
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(uploaded_file)
        return "\n".join([p.text for p in doc.paragraphs])
    else:
        return uploaded_file.read().decode("utf-8")

# Inputs
job_desc = st.text_area("Paste Job Description here:", height=200)
resume_file = st.file_uploader("Upload Resume (PDF / DOCX)", type=["pdf", "docx"])

if st.button("Check ATS Score"):
    if not job_desc or not resume_file:
        st.warning("Pehle Job Description aur Resume dono do.")
    else:
        with st.spinner("AI Check kar raha hai..."):
            resume_text = extract_text(resume_file)
            
            model = genai.GenerativeModel("gemini-1.5-flash")
            
            prompt = f"""
            You are an expert ATS (Applicant Tracking System).
            Compare this Resume with the Job Description.
            
            Job Description:
            {job_desc}
            
            Resume:
            {resume_text}
            
            Give output in this format:
            1. ATS Match Score in % (e.g. 75%)
            2. Missing Keywords
            3. Strengths
            4. Improvements / Suggestions to improve resume
            5. Final Verdict
            
            Keep language simple and in Roman Urdu / English mix.
            """
            
            response = model.generate_content(prompt)
            st.success("Result Ready!")
            st.markdown(response.text)